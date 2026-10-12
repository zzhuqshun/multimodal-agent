"""Independent read-back audit of physical synthetic materials and annotations.

This checks a controlled corpus, not arbitrary-document understanding or model accuracy.
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from .contracts import GroundTruth, TaskInput
from .integrity import require, verify_artifacts

LABELS = {"姓名": "name", "申请人姓名": "name", "护照号码": "passport", "护照有效期": "expiry",
          "离境日期": "departure", "计划出发日": "departure", "返程日期": "return",
          "签证有效起日": "visa_start", "签证有效止日": "visa_end", "邀请方地址": "address",
          "访问单位": "host", "访问开始日": "start", "访问结束日": "end", "联系人": "contact",
          "签发日期": "issued", "中文姓名": "alias_zh", "英文姓名": "alias_en"}
DATES = {"expiry", "departure", "return", "visa_start", "visa_end", "start", "end", "issued"}


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def safe_file(root, relative):
    require(isinstance(relative, str) and relative and "\\" not in relative, "unsafe manifest path")
    parts = relative.split("/")
    require(not Path(relative).is_absolute() and all(p not in ("", ".", "..") for p in parts), "unsafe manifest path")
    path = (root / relative).resolve(strict=True)
    require(path.is_relative_to(root.resolve()) and path.is_file(), "manifest path escapes dataset")
    return path


def canonical(label, raw):
    if label in DATES:
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return {"kind": "date", "value": datetime.strptime(raw, fmt).date().isoformat()}
            except ValueError:
                pass
        raise ValueError("unsupported source date")
    return {"kind": "text", "value": raw}


def read_visible(task, directory):
    """Read only declared reference materials. Does not load gold or annotations."""
    observations = defaultdict(list)
    by_anchor = {}
    aliases = {}
    for document in task.documents:
        if document.role != "reference":
            continue
        require(document.media_type == "text/plain", "development audit expects plain text references")
        path = safe_file(directory, document.relative_path)
        lines = path.read_text(encoding="utf-8").splitlines()
        require(lines and lines[0].startswith("SIMULATED"), "source missing simulation label")
        for line_number, line in enumerate(lines[1:], 2):
            label, raw = line.split(": ", 1)
            require(label in LABELS, "unknown synthetic source label")
            key = LABELS[label]
            doc_type = document.document_id.split(".", 1)[1]
            if key in {"host", "start", "end", "contact", "issued"}:
                row = 2 if doc_type == "invitation-2" else 1
                key = f"visit{row}_{key}"
            value = canonical(LABELS[label], raw)
            anchor = (document.document_id, line_number)
            observations[key].append((value, anchor))
            by_anchor[anchor] = (key, value, raw)
            if LABELS[label] in ("alias_zh", "alias_en"):
                aliases[LABELS[label]] = raw
    if aliases:
        require(set(aliases) == {"alias_zh", "alias_en"}, "incomplete visible alias mapping")
        for i, (value, anchor) in enumerate(observations["name"]):
            if value["value"] == aliases["alias_zh"]:
                value = {"kind": "text", "value": aliases["alias_en"]}
                observations["name"][i] = (value, anchor)
                key, _, raw = by_anchor[anchor]
                by_anchor[anchor] = (key, value, raw)
    return observations, by_anchor


def audit_gold(task, truth, directory):
    observations, anchors = read_visible(task, directory)
    expected = {f.gold_field_id.removeprefix("gold-"): f for f in truth.expected_fields}
    require(set(expected) == set(f.field_id.removeprefix("gold-") for f in truth.expected_form.fields), "gold coverage mismatch")
    failed_coverage = False
    for key in ("departure", "return", "visa_start", "visa_end"):
        if len(observations[key]) != 1:
            break
    else:
        departure, returned, start, end = [date.fromisoformat(observations[k][0][0]["value"])
                                           for k in ("departure", "return", "visa_start", "visa_end")]
        failed_coverage = not (start <= departure <= returned <= end)

    for key, field in expected.items():
        visible = observations[key]
        distinct = {json.dumps(v, sort_keys=True) for v, _ in visible}
        status = "missing" if not visible else ("conflict" if len(distinct) > 1 else "verified")
        issue = "missing_source" if status == "missing" else "source_conflict"
        values = [visible[0][0]] if status == "verified" else [None]
        if key == "days":
            if not observations["departure"] or not observations["return"]:
                status, values, issue = "missing", [None], "missing_calculation_input"
            else:
                departure = date.fromisoformat(observations["departure"][0][0]["value"])
                returned = date.fromisoformat(observations["return"][0][0]["value"])
                status, values = "verified", [{"kind": "integer", "value": (returned - departure).days + 1}]
        if failed_coverage and key in ("return", "days"):
            status, values, issue = ("conflict" if key == "return" else "review_required"), [None], (
                "visa_range_mismatch" if key == "return" else "upstream_conflict")
        require(field.expected_status == status, f"gold status differs from visible materials: {key}")
        actual = [v.model_dump(mode="json", exclude={"schema_version"}) if v is not None else None
                  for v in field.accepted_values]
        require(actual == values, f"gold value differs from visible materials: {key}")
        require(field.expected_issue_codes == ([] if status == "verified" else [issue]), "gold issue mismatch")
        if field.expected_calculation is not None:
            require(key == "days" and len(field.expected_calculation.inputs) == 2,
                    "unexpected synthetic calculation")
            for source_key, candidate in zip(("departure", "return"), field.expected_calculation.inputs):
                actual_input = candidate.value.model_dump(mode="json", exclude={"schema_version"})
                supported = [anchors[(a.document_id, a.locator.line_start)] for a in candidate.evidence]
                require(any(k == source_key and v == actual_input for k, v, _ in supported),
                        "calculation input differs from physical source")
        represented = []
        for candidate in field.expected_candidates:
            source_matches = []
            for anchor in candidate.evidence:
                require(anchor.locator.kind == "text" and anchor.locator.line_start == anchor.locator.line_end,
                        "synthetic evidence must anchor one physical source line")
                source_key, source_value, _ = anchors[(anchor.document_id, anchor.locator.line_start)]
                if source_key == key:
                    source_matches.append(source_value)
            value = candidate.value.model_dump(mode="json", exclude={"schema_version"})
            require(value in source_matches, f"candidate normalization/source mismatch: {key}")
            represented.append(json.dumps(value, sort_keys=True))
        require(Counter(represented) == Counter(json.dumps(v, sort_keys=True) for v, _ in visible),
                f"gold did not retain all visible candidates: {key}")
        for group in field.evidence_alternatives:
            supported = set()
            for anchor in group:
                source_key, _, _ = anchors[(anchor.document_id, anchor.locator.line_start)]
                supported.add(source_key)
            needed = {"departure", "return"} if key == "days" and status == "verified" else {key}
            if status == "verified":
                require(needed <= supported, "gold evidence alternative lacks required sources")


def audit_form(task, truth, directory):
    """Read actual field regions rather than trusting renderer-returned coordinates."""
    target = next(d for d in task.documents if d.document_id == task.target_document_id)
    path = safe_file(directory, target.relative_path)
    if target.media_type == "application/pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            require(len(pdf.pages) == 1, "template unexpectedly spans multiple pages")
            require("SIMULATED" in (pdf.pages[0].extract_text() or ""), "PDF lacks simulation marker")
            for field in truth.expected_form.fields:
                loc = field.target_locator
                page = pdf.pages[loc.page - 1]
                x0, top, x1, bottom = loc.bbox
                text = page.crop((x0 * page.width, top * page.height, x1 * page.width, bottom * page.height)).extract_text() or ""
                require(field.label in text, f"PDF label outside gold field region: {field.field_id}")
                # Target field region has label plus blank entry line, never an answer.
                compact = "".join(text.split()).replace("项目1", "").replace("项目2", "")
                require(compact == "".join(field.label.split()), "unexpected content in blank PDF field")
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        require(len(workbook.sheetnames) == 1, "unexpected or hidden answer sheet")
        worksheet = workbook.active
        require("SIMULATED" in worksheet["A2"].value, "Excel lacks simulation marker")
        for field in truth.expected_form.fields:
            loc = field.target_locator
            require(loc.sheet == worksheet.title, "unknown gold worksheet")
            cells = [cell for row in worksheet[loc.cell_range] for cell in row]
            populated = [c.value for c in cells if c.value is not None]
            require(populated == [field.label], "Excel target label/blank cells mismatch")
        require(not any(c.data_type == "f" for row in worksheet for c in row), "blank form must not contain answer formulas")
        workbook.close()


def audit_dataset(root: Path) -> dict:
    root = root.resolve(strict=True)
    manifest = read_json(root / "manifest.json")
    require(manifest["dataset_version"] == "0.1.0" and manifest["split"] == "dev", "unsupported development manifest")
    records = manifest["tasks"]
    require(len(records) == 15, "development corpus must contain 15 tasks")
    require({r["task_id"] for r in records} == {f"DEV-{n:03d}" for n in range(1, 16)}, "unexpected task IDs")
    require(Counter(r["category"] for r in records) == {"normal": 6, "conflict": 5, "missing": 4}, "category counts mismatch")
    require(Counter(r["template_family_id"] for r in records) == {"A": 5, "B": 5, "C": 5}, "template counts mismatch")
    allowed_files = {"manifest.json"}
    for relative, sha in manifest["templates"].items():
        p = safe_file(root, relative)
        require(hashlib.sha256(p.read_bytes()).hexdigest() == sha, "template hash mismatch")
        allowed_files.add(relative)
    entities, total_fields, statuses = set(), 0, Counter()
    for record in records:
        require(record["split"] == "dev" and not record["human_reviewed"], "incorrect split/review claims")
        require(not entities.intersection(record["entity_ids"]), "entity reused across development tasks")
        entities.update(record["entity_ids"])
        task_dir = (root / record["task_path"]).resolve(strict=True)
        require(task_dir.is_relative_to(root), "task directory escapes dataset")
        require({"input.json", "ground_truth.json", "annotations.json"} <= set(record["files"]),
                "task metadata files must be indexed")
        require(record["task_path"] == "tasks/" + record["task_id"], "unexpected task directory")
        for relative, sha in record["files"].items():
            p = safe_file(task_dir, relative)
            require(hashlib.sha256(p.read_bytes()).hexdigest() == sha, f"dataset artifact hash mismatch: {record['task_id']}/{relative}")
            allowed_files.add(record["task_path"] + "/" + relative)
        task = TaskInput.model_validate_json((task_dir / "input.json").read_bytes())
        truth = GroundTruth.model_validate_json((task_dir / "ground_truth.json").read_bytes())
        require(task.task_id == record["task_id"], "task metadata mismatch")
        require(all(d.synthetic for d in task.documents), "development materials must be synthetic")
        require(not truth.annotation_provenance.independently_reviewed, "synthetic gold is not independently human-reviewed")
        require(all(d.relative_path.startswith("documents/") for d in task.documents), "answer leakage through declared files")
        require({d.relative_path for d in task.documents} == {str(p.relative_to(task_dir)) for p in (task_dir / "documents").iterdir()},
                "undeclared source material")
        suffix = ".xlsx" if record["template_family_id"] == "B" else ".pdf"
        template = root / "templates" / (record["template_family_id"] + suffix)
        target = next(d for d in task.documents if d.role == "target_form")
        require((task_dir / target.relative_path).read_bytes() == template.read_bytes(), "target differs from family template")
        verify_artifacts(task, task_dir, truth=truth)
        audit_form(task, truth, task_dir)
        audit_gold(task, truth, task_dir)
        annotations = read_json(task_dir / "annotations.json")
        require(annotations["task_id"] == task.task_id and annotations["evaluation_only"] is True, "incorrect annotation metadata")
        expected_affected = [{"key": f.gold_field_id.removeprefix("gold-"), "status": f.expected_status,
                             "action": f.expected_action} for f in truth.expected_fields if f.expected_status != "verified"]
        require(annotations["expected_affected_fields"] == expected_affected, "annotation affected fields mismatch")
        for mutation in annotations["mutations"]:
            doc = next(d for d in task.documents if d.document_id == task.task_id + "." + mutation["document"])
            text = (task_dir / doc.relative_path).read_text()
            if mutation["after"] is not None:
                require(mutation["after"] in text, "mutation not present in visible materials")
            else:
                require(mutation["before"] not in text, "removed value remains visible in its source")
        require(record["anomaly_tags"] == [m["key"] for m in annotations["mutations"]], "anomaly tags mismatch")
        require(bool(expected_affected) == (record["category"] != "normal"), "task category contradicts gold")
        statuses.update(f.expected_status for f in truth.expected_fields)
        total_fields += len(truth.expected_fields)
    physical = {str(p.relative_to(root)) for p in root.rglob("*") if p.is_file()}
    require(physical == allowed_files, "unindexed or missing dataset files")
    return {"tasks": len(records), "templates": 3, "categories": dict(Counter(r["category"] for r in records)),
            "fields": total_fields, "expected_field_statuses": dict(statuses), "artifact_audit": "passed",
            "human_reviewed_tasks": 0, "model_runs": 0}
