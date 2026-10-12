"""Reproducible, fictional development tasks. No prediction or model invocation."""
from __future__ import annotations

import hashlib
import io
import json
import random
import shutil
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from .contracts import GroundTruth, TaskInput
from .integrity import require

VERSION = "0.1.0"
DEFAULT_SEED = 20261012
DEFAULT_FONT = Path("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BASE_FIELDS = [
    ("name", "姓名", "text"), ("passport", "护照号码", "text"),
    ("expiry", "护照有效期", "date"), ("departure", "离境日期", "date"),
    ("return", "返程日期", "date"), ("days", "出访天数", "integer"),
    ("visa_start", "签证有效起日", "date"), ("visa_end", "签证有效止日", "date"),
    ("address", "邀请方地址", "text"),
]
SCENARIOS = [
    ("A", "normal", "单人单次出访"),
    ("B", "normal", "无歧义的不同日期格式"),
    ("C", "normal", "两个访问项目独立关联"),
    ("A", "normal", "明确中英文姓名对应关系"),
    ("B", "normal", "离境日期与访问开始日、签发日区分"),
    ("C", "normal", "按模拟规则计算含首尾天数"),
    ("A", "conflict", "两份材料护照号不同"),
    ("B", "conflict", "两份材料护照有效期不同"),
    ("C", "conflict", "第一访问项目开始日冲突"),
    ("A", "conflict", "申请人声明与护照姓名不同"),
    ("B", "conflict", "签证区间未覆盖返程日期"),
    ("C", "missing", "签证有效期缺失"),
    ("A", "missing", "邀请方地址缺失"),
    ("B", "missing", "返程日期缺失并影响天数计算"),
    ("C", "missing", "第二访问项目联系人缺失"),
]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def typed(kind, value):
    return {"kind": kind, "value": value}


def specs(family):
    labels = {"B": {"name": "申请人姓名", "departure": "拟离境日期", "return": "计划返程日"},
              "C": {"name": "出访人员", "departure": "计划出发日", "return": "返回日期"}}
    fields = [{"key": k, "label": labels.get(family, {}).get(k, label), "kind": kind}
              for k, label, kind in BASE_FIELDS]
    if family == "C":
        for row in (1, 2):
            for key, label, kind in [("host", "访问单位", "text"), ("start", "访问开始日", "date"),
                                     ("end", "访问结束日", "date"), ("contact", "联系人", "text")]:
                fields.append({"key": f"visit{row}_{key}", "label": label, "kind": kind,
                               "group_id": "visits", "row_id": f"visit-{row}"})
    return fields


def render_form(path: Path, family: str, font: Path) -> list[dict]:
    """Return gold field regions; the blank form contains labels, never semantic keys."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    fields = specs(family)
    if family == "B":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        wb = Workbook()
        wb.properties.creator = "SmartFill synthetic generator"
        wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
        ws = wb.active
        ws.title = "申请表"
        ws.row_dimensions[1].height = 36
        ws.row_dimensions[2].height = 24
        ws.row_dimensions[3].height = 12
        ws.merge_cells("A1:D1")
        ws["A1"] = "模拟出访申请表"
        ws["A1"].font = Font(name="Droid Sans Fallback", size=18, bold=True, color="17365D")
        ws.merge_cells("A2:D2")
        ws["A2"] = "SIMULATED / 虚构实验模板，请根据提供的材料填写。"
        ws["A2"].font = Font(size=10, color="666666")
        thin = Side(style="thin", color="CCD6E0")
        row = 4
        for index, field in enumerate(fields):
            if index in (0, 3, 6):
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                ws.cell(row, 1, ["人员信息", "行程信息", "参考信息"][index // 3])
                ws.cell(row, 1).fill = PatternFill("solid", fgColor="DDE8F3")
                ws.cell(row, 1).font = Font(bold=True, color="17365D", size=12)
                ws.row_dimensions[row].height = 27
                row += 1
            ws.cell(row, 1, field["label"])
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            for cell in ws[row]:
                cell.border = Border(bottom=thin, top=thin, left=thin, right=thin)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.font = Font(name="Droid Sans Fallback", size=11)
            ws.row_dimensions[row].height = 32
            ws.cell(row, 2).number_format = {"date": "yyyy-mm-dd", "integer": "0"}.get(field["kind"], "@")
            field["locator"] = {"kind": "cell", "sheet": ws.title, "cell_range": f"A{row}:D{row}"}
            row += 1
        ws.column_dimensions["A"].width = 23
        for column in "BCD":
            ws.column_dimensions[column].width = 18
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "B5"
        ws.print_options.horizontalCentered = True
        ws.print_area = f"A1:D{row - 1}"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        stream = io.BytesIO()
        wb.save(stream)
        # ZIP member timestamps and core modified time otherwise break byte reproducibility.
        import re
        with ZipFile(stream) as source, ZipFile(path, "w", compression=ZIP_DEFLATED) as out:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    data = re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                                  rb"\g<1>2026-01-01T00:00:00Z\g<2>", data)
                member = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
                member.compress_type = ZIP_DEFLATED
                member.external_attr = 0o600 << 16
                out.writestr(member, data)
        return fields

    require(font.is_file(), "Chinese TTF font missing; supply --font")
    pdfmetrics.registerFont(TTFont("SmartFillCJK", str(font)))
    width, height = landscape(A4) if family == "C" else A4
    canvas = Canvas(str(path), pagesize=(width, height), invariant=1, pageCompression=1)
    def mixed_text(x, y, text, size):
        # Droid fallback has Chinese glyphs but no Latin glyphs.
        for character in text:
            name = "Helvetica" if ord(character) < 128 else "SmartFillCJK"
            canvas.setFont(name, size)
            canvas.drawString(x, y, character)
            x += pdfmetrics.stringWidth(character, name, size)

    canvas.setTitle("模拟出访申请表")
    canvas.setAuthor("SmartFill synthetic generator")
    canvas.setFillColor(colors.HexColor("#17365D"))
    canvas.setFont("SmartFillCJK", 21)
    canvas.drawString(40, height - 55, "模拟出访申请表")
    canvas.setFont("SmartFillCJK", 10)
    canvas.setFillColor(colors.HexColor("#666666"))
    mixed_text(40, height - 78, "SIMULATED / 虚构实验模板，请根据提供的材料填写。", 10)

    def field_box(field, x, top, w, h):
        y = height - top - h
        canvas.setStrokeColor(colors.HexColor("#BBC9D7"))
        canvas.setFillColor(colors.white)
        canvas.rect(x, y, w, h, fill=1, stroke=1)
        canvas.setFillColor(colors.HexColor("#17365D"))
        canvas.setFont("SmartFillCJK", 10)
        canvas.drawString(x + 9, y + h - 17, field["label"])
        canvas.setStrokeColor(colors.HexColor("#D3DBE3"))
        canvas.line(x + 9, y + 9, x + w - 9, y + 9)
        field["locator"] = {"kind": "page", "page": 1,
                            "bbox": [x / width, top / height, (x + w) / width, (top + h) / height]}

    if family == "A":
        for i, field in enumerate(fields):
            field_box(field, 40, 108 + i * 59, width - 80, 53)
    else:
        for i, field in enumerate(fields[:9]):
            field_box(field, 40 + (i % 3) * 254, 103 + (i // 3) * 57, 246, 51)
        canvas.setFont("SmartFillCJK", 13)
        canvas.setFillColor(colors.HexColor("#17365D"))
        canvas.drawString(40, height - 300, "访问安排（按项目编号分别填写）")
        for i, field in enumerate(fields[9:]):
            row, col = divmod(i, 4)
            field_box(field, 40 + col * 190, 315 + row * 82, 184, 72)
            if col == 0:
                canvas.setFont("SmartFillCJK", 9)
                mixed_text(49, height - (315 + row * 82) - 34, f"项目 {row + 1}", 9)
    canvas.setFont("SmartFillCJK", 9)
    canvas.setFillColor(colors.HexColor("#666666"))
    mixed_text(40, 30, "实验用空白表单 / 非真实申请材料", 9)
    canvas.save()
    return fields


def make_sources(number: int, rng: random.Random):
    """Choose fictional facts, then mutate the visible materials before gold authoring."""
    departure = date(2026, 11, 1) + timedelta(days=rng.randrange(25))
    returned = departure + timedelta(days=rng.randrange(5, 10))
    english = f"Test Person {number:02d}"
    passport = f"TEST-P{rng.randrange(10000000, 99999999)}-{number:02d}"
    docs = defaultdict(list)
    mutations = []

    def add(doc, key, label, value, kind="text", raw=None):
        canonical = value.isoformat() if isinstance(value, date) else value
        docs[doc].append({"key": key, "label": label, "raw": raw or str(canonical),
                          "value": typed(kind, canonical)})

    add("passport", "name", "姓名", english)
    add("passport", "passport", "护照号码", passport)
    add("passport", "expiry", "护照有效期", date(2031, 5, number), "date")
    add("itinerary", "departure", "离境日期", departure, "date",
        departure.strftime("%Y/%m/%d") if number == 2 else None)
    add("itinerary", "return", "返程日期", returned, "date")
    add("visa", "visa_start", "签证有效起日", departure - timedelta(days=2), "date")
    add("visa", "visa_end", "签证有效止日", returned + timedelta(days=2), "date")
    for row in (1, 2):
        doc = f"invitation-{row}"
        start = departure + timedelta(days=row)
        add(doc, "passport", "护照号码", passport)
        add(doc, f"visit{row}_host", "访问单位", f"Fictional Research Unit {number:02d}-{row}")
        add(doc, f"visit{row}_start", "访问开始日", start, "date")
        add(doc, f"visit{row}_end", "访问结束日", start + timedelta(days=1), "date")
        add(doc, f"visit{row}_contact", "联系人", f"Test Contact {number:02d}-{row}")
        add(doc, f"visit{row}_issued", "签发日期", departure - timedelta(days=10), "date")
        if row == 1:
            add(doc, "address", "邀请方地址", f"Fictional Road {number}, Test City")
    if number == 2:
        add("invitation-1", "departure", "计划出发日", departure, "date")
    if number == 4:
        add("invitation-1", "name", "姓名", english, raw=f"测试人员{number:02d}")
        add("aliases", "alias_zh", "中文姓名", f"测试人员{number:02d}")
        add("aliases", "alias_en", "英文姓名", english)
    if number == 10:
        add("applicant", "name", "申请人姓名", english)

    def change(doc, key, after):
        item = next(o for o in docs[doc] if o["key"] == key)
        before = item["raw"]
        item["raw"] = after
        item["value"] = typed(item["value"]["kind"], after)
        mutations.append({"document": doc, "key": key, "before": before, "after": after})

    def remove(doc, key):
        item = next(o for o in docs[doc] if o["key"] == key)
        docs[doc].remove(item)
        mutations.append({"document": doc, "key": key, "before": item["raw"], "after": None})

    if number == 7:
        change("invitation-1", "passport", f"TEST-CONFLICT-{number:02d}")
    elif number == 8:
        add("passport-addendum", "expiry", "护照有效期", date(2032, 5, number), "date")
        mutations.append({"document": "passport-addendum", "key": "expiry",
                          "before": date(2031, 5, number).isoformat(), "after": date(2032, 5, number).isoformat()})
    elif number == 9:
        add("invitation-update", "visit1_start", "访问开始日", departure + timedelta(days=3), "date")
        mutations.append({"document": "invitation-update", "key": "visit1_start",
                          "before": (departure + timedelta(days=1)).isoformat(),
                          "after": (departure + timedelta(days=3)).isoformat()})
    elif number == 10:
        change("applicant", "name", f"Different Test Person {number:02d}")
    elif number == 11:
        change("visa", "visa_end", (returned - timedelta(days=2)).isoformat())
    elif number == 12:
        remove("visa", "visa_start")
        remove("visa", "visa_end")
    elif number == 13:
        remove("invitation-1", "address")
    elif number == 14:
        remove("itinerary", "return")
    elif number == 15:
        remove("invitation-2", "visit2_contact")
    return docs, mutations


def build_task(root, number, family, category, title, fields, template, rng):
    task_id = f"DEV-{number:03d}"
    directory = root / "tasks" / task_id
    materials = directory / "documents"
    materials.mkdir(parents=True)
    suffix = ".xlsx" if family == "B" else ".pdf"
    shutil.copyfile(template, materials / ("form" + suffix))
    documents = []
    observations = defaultdict(list)
    sources, mutations = make_sources(number, rng)
    for doc, rows in sources.items():
        header = "SIMULATED / 虚构实验材料，非真实证件或业务文件。"
        if doc.startswith("invitation-"):
            header += " 访问项目编号：" + ("2" if doc == "invitation-2" else "1") + "。"
        if doc == "aliases":
            header += " 姓名按本文件对应关系识别，填报统一采用英文姓名。"
        lines = [header]
        document_id = f"{task_id}.{doc}"
        for item in rows:
            line = f'{item["label"]}: {item["raw"]}'
            lines.append(line)
            anchor = {"document_id": document_id, "locator": {"kind": "text", "line_start": len(lines),
                      "line_end": len(lines)}, "quote": item["raw"],
                      "start": len(item["label"]) + 2, "end": len(line)}
            observations[item["key"]].append({"value": item["value"], "evidence": [anchor]})
        (materials / f"{doc}.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    # Alias normalization is explicitly supported by visible bilingual mapping anchors.
    if number == 4:
        mapping = [observations["alias_zh"][0]["evidence"][0], observations["alias_en"][0]["evidence"][0]]
        observations["name"][-1]["evidence"].extend(mapping)

    for path in sorted(materials.iterdir()):
        target = path.stem == "form"
        documents.append({"document_id": f"{task_id}.{path.stem}", "role": "target_form" if target else "reference",
                          "media_type": (XLSX if family == "B" else "application/pdf") if target else "text/plain",
                          "relative_path": str(path.relative_to(directory)), "sha256": digest(path),
                          "synthetic": True})
    rules = [
        {"rule_id": "sim-days", "version": "1", "kind": "date_days_inclusive",
         "description": "模拟规则：出访天数=返程日期-离境日期+1；缺失或冲突输入不可计算。"},
        {"rule_id": "sim-consistency", "version": "1", "kind": "consistency",
         "description": "模拟规则：同一人员同一字段出现不兼容来源，拒填并请求确认；不默认认定较新文件优先。"},
        {"rule_id": "sim-visa-coverage", "version": "1", "kind": "consistency",
         "description": "模拟规则：签证区间已知时须覆盖行程；返程超界则返程conflict、天数review_required。签证区间缺失只标记对应签证字段missing，行程字段保留可定位值。"},
    ]
    task = {"task_id": task_id, "description": "本任务所有材料针对同一出访人员。根据提供的可见材料填写空白出访申请表。"
            "访问项目按项目编号关联；日期使用ISO格式，姓名别名仅按明确对应关系处理。"
            "所有资料和规则均为模拟；信息缺失或冲突时留空并提出对应请求。",
            "target_document_id": f"{task_id}.form", "documents": documents,
            "ruleset_id": "sim-travel", "ruleset_version": "1", "rules": rules,
            "as_of_date": "2026-10-12"}
    form_fields, expected = [], []
    for spec in fields:
        key = spec["key"]
        fid = "gold-" + key
        form_fields.append({"field_id": fid, "label": spec["label"], "value_type": spec["kind"], "required": True,
                            "target_locator": spec["locator"], "group_id": spec.get("group_id"),
                            "row_id": spec.get("row_id"), "rule_ids": (["sim-days", "sim-visa-coverage"] if key == "days" else
                            ["sim-consistency", "sim-visa-coverage"] if key == "return" else ["sim-consistency"])})
        candidates = observations[key]
        unique = {json.dumps(c["value"], sort_keys=True) for c in candidates}
        status = "missing" if not candidates else ("conflict" if len(unique) > 1 else "verified")
        code = "missing_source" if status == "missing" else "source_conflict"
        calculation = None
        alternatives = [c["evidence"] for c in candidates]
        if key == "days":
            inputs = [observations[k][0] for k in ("departure", "return") if observations[k]]
            if len(inputs) == 2:
                days = (date.fromisoformat(inputs[1]["value"]["value"]) - date.fromisoformat(inputs[0]["value"]["value"])).days + 1
                value = typed("integer", days)
                calculation = {"rule_id": "sim-days", "rule_version": "1", "inputs": inputs,
                               "expected_value": value, "unit": "days"}
                status = "verified"
                alternatives = [[a for c in inputs for a in c["evidence"]]]
            else:
                value, status, code = None, "missing", "missing_calculation_input"
        else:
            value = candidates[0]["value"] if status == "verified" else None
        if number == 11 and key in ("return", "days"):
            status = "conflict" if key == "return" else "review_required"
            value, code = None, "visa_range_mismatch" if key == "return" else "upstream_conflict"
            coverage = [observations[k][0]["evidence"][0] for k in ("departure", "return", "visa_start", "visa_end")]
            alternatives = [coverage]
            calculation = None
        expected.append({"gold_field_id": fid, "accepted_values": [value],
                         "expected_status": status, "expected_action": {"verified": "fill", "conflict": "request_confirmation",
                         "missing": "request_documents", "review_required": "request_confirmation"}[status],
                         "expected_candidates": candidates, "evidence_alternatives": alternatives,
                         "expected_issue_codes": [] if status == "verified" else [code],
                         "expected_calculation": calculation})
    truth = {"task_id": task_id, "dataset_version": VERSION,
             "expected_form": {"form_id": f"gold-{family}", "target_document_id": task["target_document_id"],
                               "fields": form_fields, "groups": [{"group_id": "visits", "label": "访问安排",
                               "repeatable": True}] if family == "C" else []},
             "expected_fields": expected,
             "annotation_provenance": {"method": "synthetic_generator", "version": VERSION, "independently_reviewed": False}}
    # Validate contracts without passing gold to any prediction code.
    TaskInput.model_validate_json(json.dumps(task))
    GroundTruth.model_validate_json(json.dumps(truth))
    write_json(directory / "input.json", task)
    write_json(directory / "ground_truth.json", truth)
    write_json(directory / "annotations.json", {"task_id": task_id, "evaluation_only": True,
               "mutations": mutations, "expected_affected_fields": [
                   {"key": f["gold_field_id"].removeprefix("gold-"), "status": f["expected_status"],
                    "action": f["expected_action"]} for f in expected if f["expected_status"] != "verified"]})
    return {"task_id": task_id, "task_path": f"tasks/{task_id}", "split": "dev", "category": category,
            "scenario": title, "template_family_id": family,
            "entity_ids": [f"fictional-person-{number:03d}"] + [f"fictional-host-{number:03d}-{i}" for i in (1, 2)],
            "ruleset_version": "1", "anomaly_tags": [m["key"] for m in mutations],
            "human_reviewed": False,
            "files": {str(p.relative_to(directory)): digest(p) for p in sorted(directory.rglob("*")) if p.is_file()}}


def generate_dataset(output: Path, seed: int = DEFAULT_SEED, font: Path = DEFAULT_FONT) -> dict:
    """Build into a sibling staging directory, audit, then rename; never overwrite data."""
    from .dataset_audit import audit_dataset
    output = output.absolute()
    require(not output.exists(), "output already exists; choose a new directory")
    require(font.is_file(), "Chinese TTF font missing; supply --font")
    output.parent.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    with tempfile.TemporaryDirectory(prefix=".smartfill-build-", dir=output.parent) as temporary:
        root = Path(temporary) / "bundle"
        templates = root / "templates"
        templates.mkdir(parents=True)
        rendered = {}
        for family in "ABC":
            path = templates / (family + (".xlsx" if family == "B" else ".pdf"))
            rendered[family] = (path, render_form(path, family, font))
        records = [build_task(root, n, family, category, title, rendered[family][1], rendered[family][0], rng)
                   for n, (family, category, title) in enumerate(SCENARIOS, 1)]
        from importlib.metadata import version
        manifest = {"dataset_version": VERSION, "generator_version": VERSION, "seed": seed, "split": "dev",
                    "source": "fictional_synthetic", "human_reviewed": False,
                    "render_environment": {"font_sha256": digest(font),
                     "packages": {p: version(p) for p in ["reportlab", "openpyxl", "pdfplumber"]}},
                    "templates": {str(p.relative_to(root)): digest(p) for p in sorted(templates.iterdir())},
                    "tasks": records}
        write_json(root / "manifest.json", manifest)
        report = audit_dataset(root)
        root.rename(output)
    return {**report, "directory": str(output), "seed": seed}
