"""Read physical files and discover supported blank-field geometry, without gold."""
from __future__ import annotations

from pathlib import Path

from .contracts import TaskInput
from .integrity import require

# A declared domain vocabulary, not a template ID or gold-field map.
LABELS = {
    "姓名": ("name", "text"), "申请人姓名": ("name", "text"), "出访人员": ("name", "text"),
    "护照号码": ("passport", "text"), "护照有效期": ("expiry", "date"),
    "离境日期": ("departure", "date"), "拟离境日期": ("departure", "date"), "计划出发日": ("departure", "date"),
    "返程日期": ("return", "date"), "计划返程日": ("return", "date"), "返回日期": ("return", "date"),
    "出访天数": ("days", "integer"), "签证有效起日": ("visa_start", "date"),
    "签证有效止日": ("visa_end", "date"), "邀请方地址": ("address", "text"),
    "访问单位": ("visit.host", "text"), "访问开始日": ("visit.start", "date"),
    "访问结束日": ("visit.end", "date"), "联系人": ("visit.contact", "text"),
    "签发日期": ("issued", "date"), "中文姓名": ("alias_zh", "text"), "英文姓名": ("alias_en", "text"),
}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def read_region(document, path: Path, locator) -> str:
    if locator.kind == "text" and document.media_type == "text/plain":
        lines = path.read_text(encoding="utf-8").splitlines()
        require(locator.line_end <= len(lines), "text locator exceeds physical document")
        return "\n".join(lines[locator.line_start - 1:locator.line_end])
    if locator.kind == "page" and document.media_type == "application/pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            require(locator.page <= len(pdf.pages), "PDF page exceeds physical document")
            page = pdf.pages[locator.page - 1]
            require(not page.rotation and not locator.source_rotation_degrees
                    and locator.source_to_canonical is None, "rotated PDF regions are not supported")
            x0, y0, x1, y1 = locator.bbox
            return page.crop((x0 * page.width, y0 * page.height,
                              x1 * page.width, y1 * page.height)).extract_text() or ""
    if locator.kind == "cell" and document.media_type == XLSX:
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=False)
        try:
            require(locator.sheet in workbook.sheetnames, "unknown physical worksheet")
            region = workbook[locator.sheet][locator.cell_range]
            if not isinstance(region, tuple):
                region = ((region,),)
            return "\n".join(str(c.value) for row in region for c in row if c.value is not None)
        finally:
            workbook.close()
    raise ValueError("unsupported physical region format")


def discover_form(document, path: Path):
    """Supported layouts: explicit PDF rectangles and adjacent blank Excel cells."""
    import re
    regions = []
    if document.media_type == "application/pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                require(not page.rotation, "rotated PDF forms are not supported")
                boxes = []
                for rect in sorted(page.rects, key=lambda r: (r["top"], r["x0"])):
                    if rect["width"] < 30 or rect["height"] < 20:
                        continue
                    bounds = (rect["x0"], rect["top"], rect["x1"], rect["bottom"])
                    text = page.crop(bounds).extract_text() or ""
                    if text.strip():
                        boxes.append((rect, text))
                project_rows = [(rect["top"], match[1]) for rect, text in boxes
                                if (match := re.search(r"项目\s*(\d+)", text))]
                for rect, text in boxes:
                    label = text.splitlines()[0].strip()
                    project = next((row for top, row in project_rows if abs(top - rect["top"]) < 1), None)
                    regions.append((label, {"kind": "page", "page": page_number,
                        "bbox": [rect["x0"] / page.width, rect["top"] / page.height,
                                 rect["x1"] / page.width, rect["bottom"] / page.height]},
                        text, project))
    elif document.media_type == XLSX:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        workbook = load_workbook(path, data_only=False)
        try:
            require(len(workbook.worksheets) == 1 and workbook.active.sheet_state == "visible",
                    "only single visible-sheet forms are supported")
            sheet = workbook.active
            require(not any(c.data_type == "f" for row in sheet for c in row), "formula-bearing forms unsupported")
            for row in sheet:
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    if any(cell.coordinate in m and m.max_col > cell.column for m in sheet.merged_cells.ranges):
                        continue  # A merged section heading, not a label beside a blank entry cell.
                    entry = sheet.cell(cell.row, cell.column + 1)
                    merged = next((m for m in sheet.merged_cells.ranges if entry.coordinate in m), None)
                    if merged is None or entry.value is not None:
                        continue
                    locator = {"kind": "cell", "sheet": sheet.title,
                               "cell_range": f"{cell.coordinate}:{get_column_letter(merged.max_col)}{merged.max_row}"}
                    regions.append((cell.value.strip(), locator, cell.value, None))
        finally:
            workbook.close()
    else:
        raise ValueError("unsupported target form media type")
    require(bool(regions), "no supported blank-field regions discovered")
    fields, segments = [], []
    for index, (label, locator, text, project) in enumerate(regions, 1):
        key, kind = LABELS.get(label, (None, "text"))
        repeat = bool(key and key.startswith("visit.") and project)
        fields.append({"field_id": f"field-{index:04d}", "label": label, "value_type": kind,
                       "required": True, "target_locator": locator, "semantic_key": key,
                       "group_id": "visit-projects" if repeat else None,
                       "row_id": "project-" + project if repeat else None,
                       "rule_ids": (["sim-days", "sim-visa-coverage"] if key == "days" else
                                    ["sim-consistency", "sim-visa-coverage"] if key == "return" else ["sim-consistency"])})
        segments.append({"segment_id": f"region-{index:04d}", "text": text, "locator": locator})
    form = {"form_id": "discovered-form", "target_document_id": document.document_id,
            "fields": fields, "groups": [{"group_id": "visit-projects", "label": "访问安排",
                                         "repeatable": True}] if any(f["group_id"] for f in fields) else []}
    parsed = {"document_id": document.document_id, "parser_name": "physical-form-regions",
              "parser_version": "0.1.0", "segments": segments}
    return form, parsed


def parse_references(task: TaskInput, root: Path):
    documents = []
    for document in task.documents:
        if document.role == "target_form":
            continue
        require(document.media_type == "text/plain", "rule baseline reference parser supports UTF-8 text only")
        segments = []
        for number, line in enumerate((root / document.relative_path).read_text(encoding="utf-8").splitlines(), 1):
            if line.strip():
                segments.append({"segment_id": f"line-{number}", "text": line,
                                 "locator": {"kind": "text", "line_start": number, "line_end": number}})
        documents.append({"document_id": document.document_id, "parser_name": "utf8-lines",
                          "parser_version": "0.1.0", "segments": segments})
    return documents
